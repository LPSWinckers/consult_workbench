import json
import io
from pathlib import Path
from types import SimpleNamespace

import pytest

from app import codex, office
from app.db import database
from app.documents import apply_edits
from test_workflows import new_project, upload, office_file


def test_admin_models_persist_and_members_cannot_change_them(clients, monkeypatch):
    admin, member = clients["admin"], clients["consultant"]
    original = admin.get("/api/settings").json()["ai"]
    config = {
        "provider": "codex",
        "chat_model": "test-chat",
        "document_model": "test-doc",
        "reasoning_effort": "medium",
    }
    try:
        assert member.put("/api/settings/ai", json=config).status_code == 403
        assert (
            admin.put("/api/settings/ai", json={**config, "chat_model": "--config=bad"}).status_code
            == 422
        )
        assert admin.put("/api/settings/ai", json=config).status_code == 200
        assert member.get("/api/settings").json()["ai"] == config
        calls = []
        monkeypatch.setattr(
            codex, "generate", lambda *a, **kw: calls.append(a) or codex.Answer(answer="OK")
        )
        assert member.post("/api/settings/ai/check").status_code == 403
        assert admin.post("/api/settings/ai/check").status_code == 200
        assert len(calls) == 1
    finally:
        admin.put("/api/settings/ai", json=original)


def test_codex_uses_separate_models_and_tool_free_structured_output(monkeypatch):
    monkeypatch.setattr(codex, "binary", lambda: "codex.exe")
    monkeypatch.setattr(
        codex,
        "configuration",
        lambda: {
            "chat_model": "chat-model",
            "document_model": "document-model",
            "reasoning_effort": "high",
        },
    )
    calls = []

    def run(args, **kw):
        calls.append(args)
        assert kw["input"] == "permitted project context"
        assert "--ignore-user-config" in args and "--ephemeral" in args
        assert args[args.index("--sandbox") + 1] == "read-only"
        assert "tools.view_image=false" in args
        for flag in ("shell_tool", "unified_exec", "apps", "multi_agent", "js_repl"):
            assert args[args.index(flag) - 1] == "--disable"
        schema = json.loads(Path(args[args.index("--output-schema") + 1]).read_text())
        assert schema["additionalProperties"] is False
        Path(args[args.index("--output-last-message") + 1]).write_text('{"answer":"OK"}')
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(codex.subprocess, "run", run)
    assert codex.generate("permitted project context", codex.Answer).answer == "OK"
    codex.generate("permitted project context", codex.Answer, document=True)
    assert [a[a.index("--model") + 1] for a in calls] == ["chat-model", "document-model"]
    monkeypatch.setattr(
        codex.subprocess,
        "run",
        lambda *a, **kw: SimpleNamespace(returncode=1, stderr="private token"),
    )
    with pytest.raises(ValueError) as err:
        codex.generate("private", codex.Answer)
    assert "private token" not in str(err.value)


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_office_roundtrip_preserves_bytes_versions_and_conflicts(clients, monkeypatch, kind):
    owner, outsider = clients["eigenaar"], clients["nieuw"]
    pid = new_project(owner, "Office " + kind)
    original = office_file(kind)
    fid = upload(owner, pid, name="Werk." + kind, content=original)
    url = f"/api/files/{fid}/office"
    assert owner.post(url + "/open").status_code == 422
    monkeypatch.setenv("OFFICE_DESKTOP_ENABLED", "true")
    launches = []
    monkeypatch.setattr(office, "launch", lambda p: launches.append(p))
    assert outsider.post(url + "/open").status_code == 404
    assert owner.post(url + "/open").status_code == 200
    working = launches[0]
    assert working.read_bytes() == owner.get(f"/api/files/{fid}/download").content
    assert owner.post(url + "/import").json()["version"] == 1
    edits = {
        "docx": {"paragraph": 2, "text": "Gewijzigd in Office"},
        "pptx": {"slide": 0, "shape": 0, "text": "Gewijzigd in Office"},
        "xlsx": {"sheet": "Analyse", "cell": "B2", "value": 42},
    }
    if kind == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(original))
        workbook.active["B2"] = 42
        out = io.BytesIO()
        workbook.save(out)
        working.write_bytes(out.getvalue())
    else:
        working.write_bytes(apply_edits("Werk." + kind, original, {"edits": [edits[kind]]}))
    saved = working.read_bytes()
    assert owner.post(url + "/import").json()["version"] == 2
    assert owner.get(f"/api/files/{fid}/download").content == saved
    assert owner.get(f"/api/files/{fid}/download?version=1").content == original
    assert outsider.get(url + "/download").status_code == 404
    upload(owner, pid, name="Werk." + kind, content=original, file_id=fid, base_version=2)
    assert owner.post(url + "/import").status_code == 409
    assert owner.post(url + "/open").status_code == 409
    assert owner.get(url + "/download").content == saved
    with database() as c:
        assert c.execute("SELECT COUNT(*) FROM versions WHERE file_id=?", (fid,)).fetchone()[0] == 3


def test_chart_sheet_column_selection_and_embedded_source(clients):
    from openpyxl import Workbook

    owner = clients["eigenaar"]
    pid = new_project(owner, "Grafiekkolommen")
    book = Workbook()
    book.active.title = "Overzicht"
    book.active.append(["Proces", "Uren", "Kosten"])
    book.active.append(["Intake", 10, 200])
    book.active.append(["Advies", 20, 400])
    extra = book.create_sheet("Pilot")
    extra.append(["Proces", "Uren", "Kosten"])
    extra.append(["Pilot A", 5, 80])
    out = io.BytesIO()
    book.save(out)
    fid = upload(owner, pid, name="Kosten.xlsx", content=out.getvalue())
    path = f"/api/files/{fid}/chart"
    result = owner.get(path + "?sheet=Pilot&value_column=2").json()
    assert result["labels"] == ["Pilot A"] and result["values"] == [80]
    assert result["value_column"] == 2 and result["range"] == "A2:C2"
    assert owner.get(path + "?sheet=Onbekend").status_code == 422
    assert owner.get(path + "?value_column=-1").status_code == 422
    assert owner.get(path + ".png?sheet=Pilot&value_column=2").content.startswith(b"\x89PNG")
    result = owner.post(
        f"/api/projects/{pid}/proposals",
        json={
            "prompt": "Maak advies",
            "type": "pptx",
            "chart_file_id": fid,
            "chart_sheet": "Pilot",
            "chart_value_column": 2,
        },
    ).json()
    assert result["payload"]["chart"]["values"] == [80]
