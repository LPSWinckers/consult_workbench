import io
import math
import json
import os
import re
import secrets
import smtplib
import time
import zipfile
from collections import defaultdict, deque
from contextlib import asynccontextmanager
from datetime import date
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import quote

from fastapi import (
    Depends,
    FastAPI,
    HTTPException,
    Request,
    Response,
    UploadFile,
    File,
    Form,
)
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from openpyxl.utils import get_column_letter

from . import ai, office
from .db import (
    database,
    initialize,
    now,
    uid,
    password_hash,
    password_matches,
    audit,
    STATES,
)
from .documents import (
    extract,
    generate,
    apply_edits,
    validate_document,
    chart_png,
    SUPPORTED,
)
from .storage import storage


@asynccontextmanager
async def lifespan(app):
    initialize()
    if ai.demo_mode() and os.getenv("SEED_DEMO", "true").lower() == "true":
        seed()
    yield


app = FastAPI(title="Meridian Consulting PoC", lifespan=lifespan)
attempts = defaultdict(deque)


@app.middleware("http")
async def same_origin_mutations(request: Request, call_next):
    origin = request.headers.get("origin")
    if request.method not in {"GET", "HEAD", "OPTIONS"} and origin:
        allowed = {
            os.getenv("APP_URL", "http://localhost:3000"),
            "http://127.0.0.1:3000",
        }
        if origin not in allowed:
            from fastapi.responses import JSONResponse

            return JSONResponse({"detail": "Onbekende oorsprong"}, status_code=403)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Cache-Control"] = "no-store"
    return response


def throttle(request):
    key = request.client.host if request.client else "local"
    q = attempts[key]
    while q and q[0] < time.time() - 60:
        q.popleft()
    if len(q) >= 20:
        raise HTTPException(429, "Te veel pogingen. Probeer over een minuut opnieuw.")
    q.append(time.time())


def current_user(request: Request):
    with database() as c:
        row = c.execute(
            "SELECT u.* FROM users u JOIN sessions s ON s.user_id=u.id WHERE s.token=? AND s.expires>?",
            (request.cookies.get("session"), time.time()),
        ).fetchone()
    if not row:
        raise HTTPException(401, "Log in om verder te gaan")
    return dict(row)


def public_user(user):
    return {k: user[k] for k in ("id", "email", "name", "role", "verified")}


def admin(user):
    if user["role"] != "admin":
        raise HTTPException(403, "Alleen een beheerder mag dit doen")


def project(c, pid, user):
    row = c.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
    if not row or (
        user["role"] != "admin"
        and not c.execute(
            "SELECT 1 FROM members WHERE project_id=? AND user_id=?", (pid, user["id"])
        ).fetchone()
    ):
        raise HTTPException(404, "Project niet gevonden")
    return dict(row)


def manager(p, user):
    if user["role"] != "admin" and p["owner_id"] != user["id"]:
        raise HTTPException(403, "Alleen de projecteigenaar of beheerder mag dit doen")


def customer(c, cid, user):
    row = c.execute("SELECT * FROM customers WHERE id=?", (cid,)).fetchone()
    permitted = (
        user["role"] == "admin"
        or c.execute(
            "SELECT 1 FROM customer_access WHERE customer_id=? AND user_id=?",
            (cid, user["id"]),
        ).fetchone()
        or c.execute(
            "SELECT 1 FROM projects p JOIN members m ON m.project_id=p.id WHERE p.customer_id=? AND m.user_id=?",
            (cid, user["id"]),
        ).fetchone()
    )
    if not row or not permitted:
        raise HTTPException(404, "Klant niet gevonden")
    return dict(row)


def file_row(c, fid, user, allow_deleted=False):
    row = c.execute("SELECT * FROM files WHERE id=?", (fid,)).fetchone()
    if not row or (row["deleted"] and not allow_deleted):
        raise HTTPException(404, "Bestand niet gevonden")
    project(c, row["project_id"], user)
    return dict(row)


def blob(c, f, number=None):
    row = c.execute(
        "SELECT blob FROM versions WHERE file_id=? AND number=?",
        (f["id"], number or f["version"]),
    ).fetchone()
    if not row:
        raise HTTPException(404, "Versie niet gevonden")
    return storage.get(row["blob"])


def clean_name(name):
    name = name.strip()
    if (
        not name
        or len(name) > 160
        or name in {".", ".."}
        or re.search(r'[\\/:*?"<>|\x00-\x1f]', name)
        or name.endswith((".", " "))
    ):
        raise HTTPException(422, "Gebruik een geldige naam zonder padtekens")
    return name


def validate_dates(start, due):
    try:
        start_value = date.fromisoformat(start) if start else None
        due_value = date.fromisoformat(due) if due else None
    except ValueError as exc:
        raise HTTPException(422, "Gebruik een datum in het formaat JJJJ-MM-DD") from exc
    if start_value and due_value and due_value < start_value:
        raise HTTPException(422, "De doeldatum moet op of na de startdatum liggen")


def check_parent(c, pid, parent_id):
    if parent_id:
        parent = c.execute(
            "SELECT * FROM files WHERE id=? AND project_id=? AND kind='folder' AND deleted=0",
            (parent_id, pid),
        ).fetchone()
        if not parent:
            raise HTTPException(422, "Map niet gevonden in dit project")


def unique_name(c, pid, parent, name, exclude=""):
    if c.execute(
        "SELECT 1 FROM files WHERE project_id=? AND parent_id IS ? AND lower(name)=lower(?) AND deleted=0 AND id!=?",
        (pid, parent, name, exclude),
    ).fetchone():
        raise HTTPException(409, "Deze naam bestaat al in de map")


def save_file(c, pid, parent, name, content, user_id):
    clean_name(name)
    check_parent(c, pid, parent)
    unique_name(c, pid, parent, name)
    fid = uid()
    c.execute(
        "INSERT INTO files VALUES(?,?,?,?,?,1,0,?)",
        (fid, pid, parent, name, "file", now()),
    )
    c.execute(
        "INSERT INTO versions VALUES(?,?,?,?,?)",
        (fid, 1, storage.put(content), user_id, now()),
    )
    audit(c, pid, user_id, "Bestand toegevoegd", name)
    return fid


def company(c):
    return json.loads(
        c.execute("SELECT value FROM settings WHERE key='company'").fetchone()["value"]
    )


class Credentials(BaseModel):
    email: str = Field(max_length=254)
    password: str = Field(min_length=10, max_length=128)
    name: str = Field(default="", max_length=100)


class CustomerInput(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    industry: str = Field(default="", max_length=200)
    description: str = Field(default="", max_length=10000)
    contacts: str = Field(default="", max_length=3000)
    goals: str = Field(default="", max_length=10000)


class ProjectInput(BaseModel):
    customer_id: str
    name: str = Field(min_length=1, max_length=120)
    description: str = Field(default="", max_length=10000)
    objectives: str = Field(default="", max_length=10000)
    start_date: str = ""
    due_date: str = ""


class ProjectUpdate(BaseModel):
    state: str | None = None
    archived: bool | None = None
    name: str | None = Field(default=None, min_length=1, max_length=120)
    description: str | None = Field(default=None, max_length=10000)
    objectives: str | None = Field(default=None, max_length=10000)
    start_date: str | None = None
    due_date: str | None = None


def email_token(c, user_id, email, kind):
    token = secrets.token_urlsafe(32)
    c.execute(
        "INSERT INTO auth_tokens VALUES(?,?,?,?)",
        (token, user_id, kind, time.time() + 3600),
    )
    link = f"{os.getenv('APP_URL', 'http://localhost:3000')}/?{kind}={token}"
    host = os.getenv("SMTP_HOST")
    if host:
        message = EmailMessage()
        message["Subject"] = (
            "Meridian: bevestig je e-mailadres"
            if kind == "verify"
            else "Meridian: nieuw wachtwoord"
        )
        message["From"], message["To"] = (
            os.getenv("SMTP_FROM", "werkruimte@example.com"),
            email,
        )
        message.set_content(f"Open deze link binnen een uur:\n{link}")
        try:
            with smtplib.SMTP(host, int(os.getenv("SMTP_PORT", "587")), timeout=15) as smtp:
                smtp.starttls()
                if os.getenv("SMTP_USER"):
                    smtp.login(os.getenv("SMTP_USER"), os.getenv("SMTP_PASSWORD", ""))
                smtp.send_message(message)
        except Exception as exc:
            raise HTTPException(503, "E-mail kon niet worden verzonden") from exc
    elif not ai.demo_mode():
        raise HTTPException(503, "Configureer SMTP voor account-e-mails")
    return link if ai.demo_mode() and not host else None


@app.get("/api/health")
def health():
    return {
        "ok": True,
        "storage": "local",
        "ai": ai.codex.configuration()["provider"],
    }


@app.post("/api/auth/register")
def register(body: Credentials, request: Request):
    throttle(request)
    email = body.email.strip().lower()
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email) or not body.name.strip():
        raise HTTPException(422, "Vul een naam en geldig e-mailadres in")
    with database() as c:
        if c.execute("SELECT 1 FROM users WHERE email=?", (email,)).fetchone():
            raise HTTPException(409, "Er bestaat al een account met dit e-mailadres")
        user_id = uid()
        c.execute(
            "INSERT INTO users VALUES(?,?,?,?,?,0)",
            (
                user_id,
                email,
                body.name.strip(),
                password_hash(body.password),
                "consultant",
            ),
        )
        link = email_token(c, user_id, email, "verify")
    return {
        "message": "Bevestig je e-mailadres om in te loggen",
        "development_link": link,
    }


@app.post("/api/auth/verify")
def verify(body: dict):
    with database() as c:
        row = c.execute(
            "SELECT * FROM auth_tokens WHERE token=? AND kind='verify' AND expires>?",
            (body.get("token"), time.time()),
        ).fetchone()
        if not row:
            raise HTTPException(400, "Deze bevestigingslink is ongeldig of verlopen")
        c.execute("UPDATE users SET verified=1 WHERE id=?", (row["user_id"],))
        c.execute("DELETE FROM auth_tokens WHERE token=?", (body["token"],))
    return {"message": "E-mailadres bevestigd. Je kunt inloggen."}


@app.post("/api/auth/forgot")
def forgot(body: dict, request: Request):
    throttle(request)
    link = None
    with database() as c:
        u = c.execute(
            "SELECT * FROM users WHERE email=?",
            (str(body.get("email", "")).lower().strip(),),
        ).fetchone()
        if u:
            link = email_token(c, u["id"], u["email"], "reset")
    return {
        "message": "Als het account bestaat, ontvang je een herstellink.",
        "development_link": link,
    }


@app.post("/api/auth/reset")
def reset(body: dict, request: Request):
    throttle(request)
    password = str(body.get("password", ""))
    if not 10 <= len(password) <= 128:
        raise HTTPException(422, "Gebruik een wachtwoord van 10 tot 128 tekens")
    with database() as c:
        row = c.execute(
            "SELECT * FROM auth_tokens WHERE token=? AND kind='reset' AND expires>?",
            (body.get("token"), time.time()),
        ).fetchone()
        if not row:
            raise HTTPException(400, "Deze herstellink is ongeldig of verlopen")
        c.execute(
            "UPDATE users SET password=? WHERE id=?",
            (password_hash(password), row["user_id"]),
        )
        c.execute(
            "DELETE FROM auth_tokens WHERE user_id=? AND kind='reset'",
            (row["user_id"],),
        )
        c.execute("DELETE FROM sessions WHERE user_id=?", (row["user_id"],))
    return {"message": "Wachtwoord gewijzigd"}


@app.post("/api/auth/login")
def login(body: Credentials, request: Request, response: Response):
    throttle(request)
    with database() as c:
        row = c.execute(
            "SELECT * FROM users WHERE email=?", (body.email.lower().strip(),)
        ).fetchone()
        if not row or not password_matches(body.password, row["password"]):
            raise HTTPException(401, "E-mailadres of wachtwoord klopt niet")
        if not row["verified"]:
            raise HTTPException(403, "Bevestig eerst je e-mailadres")
        token = secrets.token_urlsafe(32)
        c.execute("DELETE FROM sessions WHERE expires<?", (time.time(),))
        c.execute(
            "INSERT INTO sessions VALUES(?,?,?)",
            (token, row["id"], time.time() + 86400),
        )
    response.set_cookie(
        "session",
        token,
        httponly=True,
        samesite="lax",
        secure=os.getenv("COOKIE_SECURE", "false").lower() == "true",
        max_age=86400,
        path="/",
    )
    return public_user(dict(row))


@app.post("/api/auth/logout")
def logout(request: Request, response: Response):
    with database() as c:
        c.execute("DELETE FROM sessions WHERE token=?", (request.cookies.get("session"),))
    response.delete_cookie("session", path="/")
    return {"ok": True}


@app.get("/api/me")
def me(user=Depends(current_user)):
    return public_user(user)


@app.get("/api/users")
def users(user=Depends(current_user)):
    with database() as c:
        return [
            dict(r)
            for r in c.execute(
                "SELECT id,name,email,role FROM users WHERE verified=1 ORDER BY name"
            )
        ]


@app.get("/api/customers")
def customers(user=Depends(current_user)):
    with database() as c:
        rows = c.execute("SELECT * FROM customers ORDER BY name").fetchall()
        result = []
        for r in rows:
            try:
                result.append(customer(c, r["id"], user))
            except HTTPException:
                pass
        return result


@app.post("/api/customers")
def create_customer(body: CustomerInput, user=Depends(current_user)):
    if not body.name.strip():
        raise HTTPException(422, "Vul een klantnaam in")
    with database() as c:
        cid = uid()
        c.execute(
            "INSERT INTO customers VALUES(?,?,?,?,?,?,?)",
            (
                cid,
                body.name,
                body.industry,
                body.description,
                body.contacts,
                body.goals,
                now(),
            ),
        )
        c.execute("INSERT INTO customer_access VALUES(?,?)", (cid, user["id"]))
    return {"id": cid}


@app.patch("/api/customers/{cid}")
def update_customer(cid: str, body: CustomerInput, user=Depends(current_user)):
    with database() as c:
        customer(c, cid, user)
        if (
            user["role"] != "admin"
            and not c.execute(
                "SELECT 1 FROM customer_access WHERE customer_id=? AND user_id=?",
                (cid, user["id"]),
            ).fetchone()
        ):
            raise HTTPException(403, "Je hebt alleen leesrechten voor deze klant")
        c.execute(
            "UPDATE customers SET name=?,industry=?,description=?,contacts=?,goals=? WHERE id=?",
            (
                body.name,
                body.industry,
                body.description,
                body.contacts,
                body.goals,
                cid,
            ),
        )
    return {"ok": True}


@app.get("/api/customers/{cid}/access")
def customer_members(cid: str, user=Depends(current_user)):
    with database() as c:
        customer(c, cid, user)
        return {
            "can_manage": user["role"] == "admin"
            or bool(
                c.execute(
                    "SELECT 1 FROM customer_access WHERE customer_id=? AND user_id=?",
                    (cid, user["id"]),
                ).fetchone()
            ),
            "users": [
                dict(r)
                for r in c.execute(
                    "SELECT u.id,u.name,u.email FROM users u JOIN customer_access a ON a.user_id=u.id WHERE a.customer_id=?",
                    (cid,),
                )
            ],
        }


@app.post("/api/customers/{cid}/access")
def grant_customer(cid: str, body: dict, user=Depends(current_user)):
    admin(user)
    with database() as c:
        customer(c, cid, user)
        if not c.execute(
            "SELECT 1 FROM users WHERE id=? AND verified=1", (body.get("user_id"),)
        ).fetchone():
            raise HTTPException(404, "Gebruiker niet gevonden")
        c.execute("INSERT OR IGNORE INTO customer_access VALUES(?,?)", (cid, body["user_id"]))
    return {"ok": True}


@app.get("/api/projects")
def projects(user=Depends(current_user)):
    with database() as c:
        rows = c.execute(
            "SELECT p.*,c.name customer_name,u.name owner_name FROM projects p JOIN customers c ON c.id=p.customer_id JOIN users u ON u.id=p.owner_id ORDER BY p.created_at DESC"
        ).fetchall()
        return [
            dict(r)
            for r in rows
            if user["role"] == "admin"
            or c.execute(
                "SELECT 1 FROM members WHERE project_id=? AND user_id=?",
                (r["id"], user["id"]),
            ).fetchone()
        ]


@app.post("/api/projects")
def create_project(body: ProjectInput, user=Depends(current_user)):
    clean_name(body.name)
    validate_dates(body.start_date, body.due_date)
    with database() as c:
        customer(c, body.customer_id, user)
        if (
            user["role"] != "admin"
            and not c.execute(
                "SELECT 1 FROM customer_access WHERE customer_id=? AND user_id=?",
                (body.customer_id, user["id"]),
            ).fetchone()
        ):
            raise HTTPException(403, "Je hebt klanttoegang nodig om een project te maken")
        pid = uid()
        c.execute(
            "INSERT INTO projects VALUES(?,?,?,?,?,?,'Concept',?,?,0,?)",
            (
                pid,
                body.customer_id,
                body.name,
                body.description,
                body.objectives,
                user["id"],
                body.start_date,
                body.due_date,
                now(),
            ),
        )
        c.execute("INSERT INTO members VALUES(?,?)", (pid, user["id"]))
        for folder in company(c)["folders"]:
            c.execute(
                "INSERT INTO files VALUES(?,?,NULL,?,'folder',1,0,?)",
                (uid(), pid, clean_name(folder), now()),
            )
        audit(c, pid, user["id"], "Project aangemaakt", body.name)
    return {"id": pid}


@app.get("/api/projects/{pid}")
def project_detail(pid: str, user=Depends(current_user)):
    with database() as c:
        p = project(c, pid, user)
        return {
            **p,
            "customer": customer(c, p["customer_id"], user),
            "can_manage": user["role"] == "admin" or p["owner_id"] == user["id"],
            "states": STATES,
            "members": [
                dict(r)
                for r in c.execute(
                    "SELECT u.id,u.name,u.email FROM users u JOIN members m ON m.user_id=u.id WHERE m.project_id=?",
                    (pid,),
                )
            ],
            "activity": [
                dict(r)
                for r in c.execute(
                    "SELECT a.*,u.name user_name FROM activity a JOIN users u ON u.id=a.user_id WHERE project_id=? ORDER BY created_at DESC LIMIT 30",
                    (pid,),
                )
            ],
        }


@app.patch("/api/projects/{pid}")
def update_project(pid: str, body: ProjectUpdate, user=Depends(current_user)):
    with database() as c:
        p = project(c, pid, user)
        manager(p, user)
        data = body.model_dump(exclude_none=True)
        if "name" in data:
            clean_name(data["name"])
        validate_dates(data.get("start_date", p["start_date"]), data.get("due_date", p["due_date"]))
        if "state" in data and data["state"] not in STATES:
            raise HTTPException(422, "Onbekende projectstatus")
        if data:
            c.execute(
                "UPDATE projects SET " + ",".join(f"{k}=?" for k in data) + " WHERE id=?",
                (*data.values(), pid),
            )
            audit(
                c,
                pid,
                user["id"],
                "Project gewijzigd",
                json.dumps({"vorige_status": p["state"], **data}, ensure_ascii=False),
            )
    return {"ok": True}


@app.get("/api/projects/{pid}/invitations")
def invitations(pid: str, user=Depends(current_user)):
    with database() as c:
        project(c, pid, user)
        return [
            dict(r)
            for r in c.execute(
                "SELECT i.*,u.name user_name,u.email,r.name requester_name FROM invitations i JOIN users u ON u.id=i.user_id JOIN users r ON r.id=i.requester_id WHERE project_id=? ORDER BY created_at DESC",
                (pid,),
            )
        ]


@app.get("/api/approvals")
def approvals(user=Depends(current_user)):
    with database() as c:
        rows = c.execute(
            "SELECT i.*,u.name user_name,p.name project_name,p.owner_id FROM invitations i JOIN users u ON u.id=i.user_id JOIN projects p ON p.id=i.project_id WHERE i.state='pending'"
        ).fetchall()
        return [dict(r) for r in rows if user["role"] == "admin" or r["owner_id"] == user["id"]]


@app.post("/api/projects/{pid}/invitations")
def request_invitation(pid: str, body: dict, user=Depends(current_user)):
    with database() as c:
        c.execute("BEGIN IMMEDIATE")
        project(c, pid, user)
        target = body.get("user_id")
        if not c.execute("SELECT 1 FROM users WHERE id=? AND verified=1", (target,)).fetchone():
            raise HTTPException(404, "Gebruiker niet gevonden")
        if (
            c.execute(
                "SELECT 1 FROM members WHERE project_id=? AND user_id=?", (pid, target)
            ).fetchone()
            or c.execute(
                "SELECT 1 FROM invitations WHERE project_id=? AND user_id=? AND state='pending'",
                (pid, target),
            ).fetchone()
        ):
            raise HTTPException(409, "Gebruiker is al lid of heeft een open uitnodiging")
        c.execute(
            "INSERT INTO invitations VALUES(?,?,?,?,'pending',NULL,?)",
            (uid(), pid, target, user["id"], now()),
        )
        audit(c, pid, user["id"], "Uitnodiging aangevraagd")
    return {"ok": True}


@app.post("/api/invitations/{iid}/decision")
def decide_invitation(iid: str, body: dict, user=Depends(current_user)):
    with database() as c:
        c.execute("BEGIN IMMEDIATE")
        row = c.execute("SELECT * FROM invitations WHERE id=?", (iid,)).fetchone()
        if not row:
            raise HTTPException(404, "Uitnodiging niet gevonden")
        manager(project(c, row["project_id"], user), user)
        if row["state"] != "pending":
            raise HTTPException(409, "Uitnodiging is al behandeld")
        state = body.get("state")
        if state not in {"approved", "rejected"}:
            raise HTTPException(422, "Ongeldige beslissing")
        c.execute(
            "UPDATE invitations SET state=?,decided_by=? WHERE id=?",
            (state, user["id"], iid),
        )
        if state == "approved":
            c.execute(
                "INSERT OR IGNORE INTO members VALUES(?,?)",
                (row["project_id"], row["user_id"]),
            )
        audit(
            c,
            row["project_id"],
            user["id"],
            "Uitnodiging " + ("goedgekeurd" if state == "approved" else "afgewezen"),
        )
    return {"ok": True}


@app.delete("/api/projects/{pid}/members/{target}")
def remove_member(pid: str, target: str, user=Depends(current_user)):
    with database() as c:
        p = project(c, pid, user)
        manager(p, user)
        if target == p["owner_id"]:
            raise HTTPException(422, "De projecteigenaar kan niet worden verwijderd")
        c.execute("DELETE FROM members WHERE project_id=? AND user_id=?", (pid, target))
        audit(c, pid, user["id"], "Projectlid verwijderd")
    return {"ok": True}


@app.get("/api/projects/{pid}/files")
def files(pid: str, user=Depends(current_user)):
    with database() as c:
        project(c, pid, user)
        return [
            dict(r)
            for r in c.execute(
                "SELECT * FROM files WHERE project_id=? ORDER BY kind DESC,name", (pid,)
            )
        ]


@app.post("/api/projects/{pid}/folders")
def create_folder(pid: str, body: dict, user=Depends(current_user)):
    with database() as c:
        project(c, pid, user)
        parent, name = body.get("parent_id"), clean_name(str(body.get("name", "")))
        check_parent(c, pid, parent)
        unique_name(c, pid, parent, name)
        fid = uid()
        c.execute(
            "INSERT INTO files VALUES(?,?,?,?,'folder',1,0,?)",
            (fid, pid, parent, name, now()),
        )
        audit(c, pid, user["id"], "Map toegevoegd", name)
    return {"id": fid}


@app.post("/api/projects/{pid}/upload")
async def upload(
    pid: str,
    file: UploadFile = File(...),
    parent_id: str = Form(""),
    file_id: str = Form(""),
    base_version: int = Form(0),
    user=Depends(current_user),
):
    with database() as c:
        project(c, pid, user)
    content = await file.read(20_000_001)
    if len(content) > 20_000_000:
        raise HTTPException(413, "Bestanden mogen maximaal 20 MB zijn")
    try:
        validate_document(file.filename or "", content)
    except Exception as exc:
        raise HTTPException(422, "Ongeldig of te groot Office-bestand") from exc
    with database() as c:
        project(c, pid, user)
        if file_id:
            f = file_row(c, file_id, user)
            if f["project_id"] != pid or f["kind"] != "file":
                raise HTTPException(422, "Ongeldig doelbestand")
            if Path(f["name"]).suffix.lower() != Path(file.filename).suffix.lower():
                raise HTTPException(422, "Het bestandstype moet gelijk blijven")
            c.execute(
                "UPDATE files SET version=version+1 WHERE id=? AND version=?",
                (file_id, base_version),
            )
            if c.execute("SELECT changes()").fetchone()[0] != 1:
                raise HTTPException(409, "Er is een nieuwere versie. Vernieuw het bestand.")
            c.execute(
                "INSERT INTO versions VALUES(?,?,?,?,?)",
                (file_id, base_version + 1, storage.put(content), user["id"], now()),
            )
            audit(c, pid, user["id"], "Nieuwe bestandsversie", f["name"])
            return {"id": file_id}
        return {"id": save_file(c, pid, parent_id or None, file.filename, content, user["id"])}


@app.patch("/api/files/{fid}")
def update_file(fid: str, body: dict, user=Depends(current_user)):
    with database() as c:
        c.execute("BEGIN IMMEDIATE")
        f = file_row(c, fid, user, True)
        if "deleted" in body:
            deleted = int(bool(body["deleted"]))
            if not deleted:
                check_parent(c, f["project_id"], f["parent_id"])
                unique_name(c, f["project_id"], f["parent_id"], f["name"], fid)
            descendants = [fid]
            for item in descendants:
                descendants.extend(
                    r["id"] for r in c.execute("SELECT id FROM files WHERE parent_id=?", (item,))
                )
            if not deleted:
                for item in descendants:
                    child = c.execute("SELECT * FROM files WHERE id=?", (item,)).fetchone()
                    unique_name(c, f["project_id"], child["parent_id"], child["name"], item)
            for item in descendants:
                c.execute("UPDATE files SET deleted=? WHERE id=?", (deleted, item))
            audit(
                c,
                f["project_id"],
                user["id"],
                "Bestand verwijderd" if deleted else "Bestand hersteld",
                f["name"],
            )
        else:
            if f["deleted"]:
                raise HTTPException(422, "Herstel het bestand eerst")
            name = clean_name(str(body.get("name", f["name"])))
            if f["kind"] == "file" and Path(name).suffix.lower() != Path(f["name"]).suffix.lower():
                raise HTTPException(422, "Behoud de bestandsextensie")
            parent = body.get("parent_id", f["parent_id"])
            check_parent(c, f["project_id"], parent)
            ancestor = parent
            while ancestor:
                if ancestor == fid:
                    raise HTTPException(422, "Een map kan niet in zichzelf worden geplaatst")
                ancestor = c.execute(
                    "SELECT parent_id FROM files WHERE id=?", (ancestor,)
                ).fetchone()["parent_id"]
            unique_name(c, f["project_id"], parent, name, fid)
            c.execute("UPDATE files SET name=?,parent_id=? WHERE id=?", (name, parent, fid))
            audit(c, f["project_id"], user["id"], "Bestand verplaatst of hernoemd", name)
    return {"ok": True}


@app.get("/api/files/{fid}/versions")
def versions(fid: str, user=Depends(current_user)):
    with database() as c:
        file_row(c, fid, user)
        return [
            dict(r)
            for r in c.execute(
                "SELECT v.number,v.created_at,u.name user_name FROM versions v JOIN users u ON u.id=v.user_id WHERE file_id=? ORDER BY number DESC",
                (fid,),
            )
        ]


def download_bytes(content, name, media="application/octet-stream"):
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
    )


@app.get("/api/files/{fid}/download")
def download(fid: str, version: int | None = None, user=Depends(current_user)):
    with database() as c:
        f = file_row(c, fid, user)
        if f["kind"] == "folder":
            return download_bytes(
                archive_content(c, f["project_id"], fid), f["name"] + ".zip", "application/zip"
            )
        return download_bytes(blob(c, f, version), f["name"])


@app.get("/api/files/{fid}/preview")
def preview(fid: str, user=Depends(current_user)):
    with database() as c:
        f = file_row(c, fid, user)
        if f["kind"] != "file":
            raise HTTPException(422, "Dit is een map")
        try:
            return {**extract(f["name"], blob(c, f)), "version": f["version"]}
        except Exception as exc:
            raise HTTPException(422, "Het bestand kan niet worden gelezen") from exc


def archive_content(c, pid, root=None):
    rows = {
        r["id"]: dict(r)
        for r in c.execute("SELECT * FROM files WHERE project_id=? AND deleted=0", (pid,))
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for f in rows.values():
            parts, parent = [f["name"]], f["parent_id"]
            belongs = root is None or f["id"] == root
            while parent in rows and f["id"] != root:
                parts.insert(0, rows[parent]["name"])
                if parent == root:
                    belongs = True
                    break
                parent = rows[parent]["parent_id"]
            if belongs:
                path = "/".join(parts)
                z.writestr(
                    path + "/" if f["kind"] == "folder" else path,
                    b"" if f["kind"] == "folder" else blob(c, f),
                )
    return out.getvalue()


@app.get("/api/projects/{pid}/download")
def download_project(pid: str, user=Depends(current_user)):
    with database() as c:
        p = project(c, pid, user)
        return download_bytes(
            archive_content(c, pid), clean_name(p["name"]) + ".zip", "application/zip"
        )


class MessageInput(BaseModel):
    body: str = Field(min_length=1, max_length=10000)


@app.get("/api/projects/{pid}/messages")
def messages(pid: str, user=Depends(current_user)):
    with database() as c:
        project(c, pid, user)
        return [
            dict(r)
            for r in c.execute(
                "SELECT m.*,u.name user_name FROM messages m JOIN users u ON u.id=m.user_id WHERE project_id=? ORDER BY created_at",
                (pid,),
            )
        ]


@app.post("/api/projects/{pid}/messages")
def post_message(pid: str, body: MessageInput, user=Depends(current_user)):
    with database() as c:
        project(c, pid, user)
        if not body.body.strip():
            raise HTTPException(422, "Schrijf een bericht")
        c.execute(
            "INSERT INTO messages VALUES(?,?,?,?,?)",
            (uid(), pid, user["id"], body.body.strip(), now()),
        )
    return {"ok": True}


@app.get("/api/projects/{pid}/conversations")
def conversations(pid: str, user=Depends(current_user)):
    with database() as c:
        project(c, pid, user)
        return [
            dict(r)
            for r in c.execute(
                "SELECT * FROM conversations WHERE project_id=? AND (user_id=? OR shared=1) ORDER BY created_at DESC",
                (pid, user["id"]),
            )
        ]


@app.post("/api/projects/{pid}/conversations")
def create_conversation(pid: str, body: dict, user=Depends(current_user)):
    with database() as c:
        project(c, pid, user)
        cid = uid()
        c.execute(
            "INSERT INTO conversations VALUES(?,?,?,?,0,?)",
            (
                cid,
                pid,
                user["id"],
                str(body.get("title", "Nieuw gesprek"))[:120],
                now(),
            ),
        )
    return {"id": cid}


def conversation(c, cid, user):
    row = c.execute("SELECT * FROM conversations WHERE id=?", (cid,)).fetchone()
    if not row:
        raise HTTPException(404, "Gesprek niet gevonden")
    project(c, row["project_id"], user)
    if row["user_id"] != user["id"] and not row["shared"]:
        raise HTTPException(404, "Gesprek niet gevonden")
    return dict(row)


@app.get("/api/conversations/{cid}")
def conversation_detail(cid: str, user=Depends(current_user)):
    with database() as c:
        row = conversation(c, cid, user)
        return {
            **row,
            "messages": [
                {**dict(r), "sources": json.loads(r["sources"])}
                for r in c.execute(
                    "SELECT * FROM ai_messages WHERE conversation_id=? ORDER BY created_at",
                    (cid,),
                )
            ],
        }


@app.patch("/api/conversations/{cid}")
def share_conversation(cid: str, body: dict, user=Depends(current_user)):
    with database() as c:
        row = conversation(c, cid, user)
        if row["user_id"] != user["id"]:
            raise HTTPException(403, "Alleen de gesprekeigenaar kan dit gesprek delen")
        c.execute(
            "UPDATE conversations SET shared=? WHERE id=?",
            (int(bool(body.get("shared"))), cid),
        )
    return {"ok": True}


def project_context(c, pid, user, selected=None):
    p = project(c, pid, user)
    cust = customer(c, p["customer_id"], user)
    context = f"Adviesbureau: {company(c)['name']}\nKlant: {cust['name']}\nSector: {cust['industry']}\nAchtergrond: {cust['description']}\nContacten: {cust['contacts']}\nKlantdoelen: {cust['goals']}\nProject: {p['name']}\nScope: {p['description']}\nDoelen: {p['objectives']}\nStatus: {p['state']}\n"
    sources = []
    for row in c.execute(
        "SELECT * FROM files WHERE project_id=? AND kind='file' AND deleted=0 ORDER BY name",
        (pid,),
    ):
        f = dict(row)
        if selected is not None and f["id"] not in selected:
            continue
        text = extract(f["name"], blob(c, f))["text"]
        available = max(0, 70000 - len(context))
        if not available:
            break
        context += (
            f"\nBestand: {f['name']} (versie {f['version']})\n" + text[: min(14000, available)]
        )
        sources.append(
            {
                "id": f["id"],
                "name": f["name"],
                "version": f["version"],
                "truncated": len(text) > min(14000, available),
            }
        )
    return context, sources


def selected_agent(c, aid, pid):
    if not aid:
        return "", None
    agent = c.execute("SELECT * FROM agents WHERE id=? AND project_id=?", (aid, pid)).fetchone()
    if not agent:
        raise HTTPException(404, "Agent niet gevonden")
    ids = json.loads(agent["file_ids"])
    return agent["instructions"], ids or None


class Question(BaseModel):
    question: str = Field(min_length=1, max_length=6000)
    agent_id: str | None = None
    file_id: str | None = None


@app.post("/api/conversations/{cid}/ask")
def ask(cid: str, body: Question, user=Depends(current_user)):
    with database() as c:
        row = conversation(c, cid, user)
        if row["user_id"] != user["id"]:
            raise HTTPException(
                403, "Dit gedeelde gesprek is alleen-lezen. Begin je eigen gesprek."
            )
        instructions, selected = selected_agent(c, body.agent_id, row["project_id"])
        if body.agent_id and "chat" not in json.loads(
            c.execute("SELECT tools FROM agents WHERE id=?", (body.agent_id,)).fetchone()["tools"]
        ):
            raise HTTPException(403, "Deze agent mag geen chat gebruiken")
        if body.file_id:
            target = file_row(c, body.file_id, user)
            if target["project_id"] != row["project_id"] or target["kind"] != "file":
                raise HTTPException(422, "Ongeldig contextbestand")
            if selected is not None and body.file_id not in selected:
                raise HTTPException(403, "Bestand valt buiten de agentreferenties")
            selected = [body.file_id]
        context, sources = project_context(c, row["project_id"], user, selected)
        available_sources = {s["id"]: s["version"] for s in sources}
        history = []
        previous = c.execute(
            "SELECT role,body,sources FROM ai_messages WHERE conversation_id=? ORDER BY created_at DESC LIMIT 12",
            (cid,),
        ).fetchall()[::-1]
        for message in previous:
            old_sources = json.loads(message["sources"])
            if all(available_sources.get(s["id"]) == s["version"] for s in old_sources):
                history.append({"role": message["role"], "content": message["body"][:4000]})
    try:
        answer = ai.answer(body.question, context, history, instructions)
    except Exception as exc:
        raise HTTPException(
            502,
            str(exc)
            if isinstance(exc, ValueError)
            else "AI-aanvraag mislukt. Controleer de Codex-configuratie en probeer opnieuw.",
        ) from exc
    with database() as c:
        conversation(c, cid, user)
        # Recheck access after the external call.
        c.execute(
            "INSERT INTO ai_messages VALUES(?,?,?,?,'[]',?)",
            (uid(), cid, "user", body.question, now()),
        )
        c.execute(
            "INSERT INTO ai_messages VALUES(?,?,?,?,?,?)",
            (uid(), cid, "assistant", answer, json.dumps(sources), now()),
        )
        if row["title"] == "Nieuw gesprek":
            c.execute("UPDATE conversations SET title=? WHERE id=?", (body.question[:70], cid))
    return {"answer": answer, "sources": sources}


class AgentInput(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    instructions: str = Field(min_length=1, max_length=10000)
    tools: list[str] = Field(default_factory=list)
    file_ids: list[str] = Field(default_factory=list)


@app.get("/api/projects/{pid}/agents")
def agents(pid: str, user=Depends(current_user)):
    with database() as c:
        project(c, pid, user)
        return [
            {
                **dict(r),
                "tools": json.loads(r["tools"]),
                "file_ids": json.loads(r["file_ids"]),
            }
            for r in c.execute("SELECT * FROM agents WHERE project_id=?", (pid,))
        ]


@app.post("/api/projects/{pid}/agents")
def create_agent(pid: str, body: AgentInput, user=Depends(current_user)):
    if set(body.tools) - {"chat", "documents", "review", "charts"}:
        raise HTTPException(422, "Onbekende agentfunctie")
    with database() as c:
        project(c, pid, user)
        for fid in body.file_ids:
            if file_row(c, fid, user)["project_id"] != pid:
                raise HTTPException(422, "Referentiebestand hoort niet bij dit project")
        aid = uid()
        c.execute(
            "INSERT INTO agents VALUES(?,?,?,?,?,?)",
            (
                aid,
                pid,
                body.name,
                body.instructions,
                json.dumps(body.tools),
                json.dumps(body.file_ids),
            ),
        )
        audit(c, pid, user["id"], "Agent aangemaakt", body.name)
    return {"id": aid}


@app.delete("/api/projects/{pid}/agents/{aid}")
def delete_agent(pid: str, aid: str, user=Depends(current_user)):
    with database() as c:
        manager(project(c, pid, user), user)
        c.execute("DELETE FROM agents WHERE id=? AND project_id=?", (aid, pid))
    return {"ok": True}


class ProposalInput(BaseModel):
    prompt: str = Field(min_length=1, max_length=6000)
    type: str = "docx"
    file_id: str | None = None
    name: str = "Adviesdocument"
    agent_id: str | None = None
    chart_file_id: str | None = None
    chart_sheet: str | None = None
    chart_label_column: int = Field(default=0, ge=0, le=1000)
    chart_value_column: int = Field(default=1, ge=0, le=1000)


@app.post("/api/projects/{pid}/proposals")
def create_proposal(pid: str, body: ProposalInput, user=Depends(current_user)):
    if body.type not in SUPPORTED:
        raise HTTPException(422, "Onbekend documenttype")
    with database() as c:
        project(c, pid, user)
        instructions, selected = selected_agent(c, body.agent_id, pid)
        if body.agent_id:
            row = c.execute("SELECT tools FROM agents WHERE id=?", (body.agent_id,)).fetchone()
            required = "review" if body.file_id else "documents"
            if required not in json.loads(row["tools"]):
                raise HTTPException(403, "Deze agent mag deze functie niet gebruiken")
            if body.chart_file_id and "charts" not in json.loads(row["tools"]):
                raise HTTPException(403, "Deze agent mag geen grafieken gebruiken")
        context, sources = project_context(c, pid, user, selected)
        f = file_row(c, body.file_id, user) if body.file_id else None
        if f and (f["project_id"] != pid or f["kind"] != "file"):
            raise HTTPException(422, "Ongeldig doelbestand")
        if f and selected is not None and f["id"] not in selected:
            raise HTTPException(403, "Dit bestand valt buiten de agentreferenties")
        name = f["name"] if f else clean_name(body.name + "." + body.type)
        if f:
            context += (
                "\nBewerk alleen dit doelbestand:\n"
                + extract(f["name"], blob(c, f))["text"][:40000]
            )
        chart = None
        if body.chart_file_id:
            if selected is not None and body.chart_file_id not in selected:
                raise HTTPException(403, "De grafiekbron valt buiten de agentreferenties")
            chart = chart_data(
                c,
                body.chart_file_id,
                user,
                body.chart_sheet,
                body.chart_label_column,
                body.chart_value_column,
            )
            if chart["project_id"] != pid:
                raise HTTPException(422, "Grafiekbron hoort niet bij dit project")
    try:
        payload = ai.propose(body.prompt, context, body.type, bool(f), instructions)
        payload.update(
            {
                "type": Path(f["name"]).suffix.lstrip(".") if f else body.type,
                "sources": sources,
                "chart": chart,
            }
        )
        # Validate the proposed edit before persisting; never execute generated code.
        if f and payload.get("edits"):
            with database() as c:
                apply_edits(f["name"], blob(c, f), payload)
    except Exception as exc:
        raise HTTPException(
            502,
            str(exc)
            if isinstance(exc, ValueError)
            else "Er kon geen geldig voorstel worden gemaakt",
        ) from exc
    with database() as c:
        project(c, pid, user)
        proposal_id = uid()
        c.execute(
            "INSERT INTO proposals VALUES(?,?,?,?,?,?,?,0,?)",
            (
                proposal_id,
                pid,
                user["id"],
                body.file_id,
                f["version"] if f else None,
                name,
                json.dumps(payload),
                now(),
            ),
        )
    return {
        "id": proposal_id,
        "name": name,
        "payload": payload,
        "base_version": f["version"] if f else None,
    }


def proposal_row(c, prid, user):
    row = c.execute(
        "SELECT * FROM proposals WHERE id=? AND user_id=?", (prid, user["id"])
    ).fetchone()
    if not row:
        raise HTTPException(404, "Voorstel niet gevonden")
    project(c, row["project_id"], user)
    return dict(row)


def proposal_content(c, row, user):
    payload = json.loads(row["payload"])
    if row["file_id"]:
        f = file_row(c, row["file_id"], user)
        if f["version"] != row["base_version"]:
            raise HTTPException(
                409, "Bestand is gewijzigd sinds dit voorstel. Maak een nieuw voorstel."
            )
        if not payload.get("edits"):
            raise HTTPException(422, "Dit voorstel bevat geen wijzigingen")
        return apply_edits(f["name"], blob(c, f), payload)
    chart = payload.get("chart")
    image = None
    if chart:
        source = file_row(c, chart["file_id"], user)
        if source["version"] != chart["version"]:
            raise HTTPException(409, "De grafiekbron is gewijzigd. Maak een nieuw voorstel.")
        image = chart_png(chart["labels"], chart["values"], chart["title"])
    return generate(payload, company(c), image)


@app.get("/api/proposals/{prid}/download")
def proposal_download(prid: str, user=Depends(current_user)):
    with database() as c:
        row = proposal_row(c, prid, user)
        return download_bytes(proposal_content(c, row, user), row["name"])


@app.post("/api/proposals/{prid}/apply")
def apply_proposal(prid: str, body: dict, user=Depends(current_user)):
    with database() as c:
        c.execute("BEGIN IMMEDIATE")
        row = proposal_row(c, prid, user)
        if row["applied"]:
            raise HTTPException(409, "Voorstel is al toegepast")
        payload = json.loads(row["payload"])
        for source in payload.get("sources", []):
            current = file_row(c, source["id"], user)
            if current["version"] != source["version"]:
                raise HTTPException(
                    409, "Een contextbestand is gewijzigd. Maak een nieuw voorstel."
                )
        content = proposal_content(c, row, user)
        if row["file_id"]:
            fid = row["file_id"]
            c.execute(
                "UPDATE files SET version=version+1 WHERE id=? AND version=?",
                (fid, row["base_version"]),
            )
            if c.execute("SELECT changes()").fetchone()[0] != 1:
                raise HTTPException(409, "Er is een nieuwere versie")
            c.execute(
                "INSERT INTO versions VALUES(?,?,?,?,?)",
                (fid, row["base_version"] + 1, storage.put(content), user["id"], now()),
            )
            audit(c, row["project_id"], user["id"], "AI-voorstel toegepast", row["name"])
        else:
            fid = save_file(
                c,
                row["project_id"],
                body.get("parent_id"),
                row["name"],
                content,
                user["id"],
            )
        c.execute("UPDATE proposals SET applied=1 WHERE id=?", (prid,))
    return {"id": fid}


def chart_data(c, fid, user, sheet=None, label_column=0, value_column=1):
    f = file_row(c, fid, user)
    if not f["name"].lower().endswith(".xlsx"):
        raise HTTPException(422, "Kies een Excel-bestand")
    sheets = extract(f["name"], blob(c, f))["sheets"]
    s = (
        next((s for s in sheets if s["name"] == sheet), None)
        if sheet
        else (sheets[0] if sheets else None)
    )
    if not s:
        raise HTTPException(422, "Werkblad niet gevonden")
    if not 0 <= label_column <= 1000 or not 0 <= value_column <= 1000:
        raise HTTPException(422, "Ongeldige grafiekkolom")
    points = [
        (str(r[label_column]), float(r[value_column]))
        for r in s["rows"][1:]
        if len(r) > max(label_column, value_column)
        and isinstance(r[value_column], (int, float))
        and not isinstance(r[value_column], bool)
        and math.isfinite(r[value_column])
        and r[label_column] is not None
    ][:50]
    if not points:
        raise HTTPException(
            422,
            "Kies een labelkolom en een kolom met numerieke waarden. Formules worden lokaal niet berekend.",
        )
    return {
        "file_id": fid,
        "project_id": f["project_id"],
        "version": f["version"],
        "sheet": s["name"],
        "title": f["name"] + " · " + s["name"],
        "range": f"{get_column_letter(label_column + 1)}2:{get_column_letter(value_column + 1)}{len(s['rows'])}",
        "label_column": label_column,
        "value_column": value_column,
        "labels": [p[0] for p in points],
        "values": [p[1] for p in points],
    }


@app.get("/api/files/{fid}/chart")
def chart(
    fid: str,
    sheet: str | None = None,
    label_column: int = 0,
    value_column: int = 1,
    user=Depends(current_user),
):
    with database() as c:
        return chart_data(c, fid, user, sheet, label_column, value_column)


@app.get("/api/files/{fid}/chart.png")
def chart_image(
    fid: str,
    sheet: str | None = None,
    label_column: int = 0,
    value_column: int = 1,
    user=Depends(current_user),
):
    with database() as c:
        data = chart_data(c, fid, user, sheet, label_column, value_column)
        return download_bytes(
            chart_png(data["labels"], data["values"], data["title"]),
            "analyse.png",
            "image/png",
        )


@app.get("/api/settings")
def settings(user=Depends(current_user)):
    with database() as c:
        return {
            "company": company(c),
            "ai_mode": ai.codex.configuration()["provider"],
            "storage": "local",
            "microsoft_connected": False,
            "office_desktop_enabled": os.getenv("OFFICE_DESKTOP_ENABLED", "false").lower()
            == "true",
            "ai": ai.codex.configuration(),
            "models": ai.codex.available_models(),
        }


@app.get("/api/files/{fid}/office")
def office_status(fid: str, user=Depends(current_user)):
    with database() as c:
        file_row(c, fid, user)
        row = c.execute(
            "SELECT * FROM office_checkouts WHERE file_id=? AND user_id=?", (fid, user["id"])
        ).fetchone()
        return {
            "enabled": office.enabled(),
            "checked_out": bool(row),
            "base_version": row["base_version"] if row else None,
        }


@app.post("/api/files/{fid}/office/open")
def open_office(fid: str, user=Depends(current_user)):
    if not office.enabled():
        raise HTTPException(422, "Schakel OFFICE_DESKTOP_ENABLED in op de werk-PC")
    with database() as c:
        c.execute("BEGIN IMMEDIATE")
        f = file_row(c, fid, user)
        if f["kind"] != "file":
            raise HTTPException(422, "Kies een Office-bestand")
        row = c.execute(
            "SELECT * FROM office_checkouts WHERE file_id=? AND user_id=?", (fid, user["id"])
        ).fetchone()
        if row and row["base_version"] != f["version"]:
            raise HTTPException(
                409,
                "De werkversie is verouderd. Download je Office-kopie en upload deze bewust als nieuwe versie.",
            )
        if not row:
            content = blob(c, f)
            row = {
                "id": uid(),
                "file_id": fid,
                "user_id": user["id"],
                "base_version": f["version"],
                "name": f["name"],
                "digest": office.digest(content),
            }
            target = office.path(row)
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(content)
            c.execute("INSERT INTO office_checkouts VALUES(?,?,?,?,?,?)", tuple(row.values()))
        target = office.path(row)
    try:
        office.launch(target)
    except OSError as exc:
        raise HTTPException(
            502, "Office kon niet openen. Controleer de installatie en standaardapps op de werk-PC."
        ) from exc
    return {"message": "Bestand geopend. Sla op en sluit Office voordat je de versie importeert."}


@app.get("/api/files/{fid}/office/download")
def download_office_copy(fid: str, user=Depends(current_user)):
    with database() as c:
        file_row(c, fid, user)
        row = c.execute(
            "SELECT * FROM office_checkouts WHERE file_id=? AND user_id=?", (fid, user["id"])
        ).fetchone()
        if not row:
            raise HTTPException(404, "Geen Office-werkkopie")
        try:
            content = office.path(row).read_bytes()
        except OSError as exc:
            raise HTTPException(
                409, "De werkkopie is niet beschikbaar. Sluit Office en probeer opnieuw."
            ) from exc
        return download_bytes(content, row["name"])


@app.post("/api/files/{fid}/office/import")
def import_office_copy(fid: str, user=Depends(current_user)):
    if not office.enabled():
        raise HTTPException(422, "Desktop Office is uitgeschakeld")
    with database() as c:
        c.execute("BEGIN IMMEDIATE")
        f = file_row(c, fid, user)
        row = c.execute(
            "SELECT * FROM office_checkouts WHERE file_id=? AND user_id=?", (fid, user["id"])
        ).fetchone()
        if not row:
            raise HTTPException(404, "Open het bestand eerst in Office")
        if row["base_version"] != f["version"]:
            raise HTTPException(
                409,
                "Een collega heeft een nieuwere versie opgeslagen. Download je werkkopie en vergelijk de wijzigingen voordat je een vervangende versie uploadt.",
            )
        try:
            with office.path(row).open("rb") as source:
                content = source.read(20_000_001)
            if len(content) > 20_000_000:
                raise ValueError("Bestand groter dan 20 MB")
            validate_document(row["name"], content)
        except Exception as exc:
            raise HTTPException(
                422, "Sla op en sluit Office. De werkkopie is niet leesbaar of ongeldig."
            ) from exc
        digest = office.digest(content)
        if digest == row["digest"]:
            return {"version": f["version"], "message": "Geen wijzigingen om op te slaan."}
        version = f["version"] + 1
        c.execute(
            "INSERT INTO versions VALUES(?,?,?,?,?)",
            (fid, version, storage.put(content), user["id"], now()),
        )
        c.execute("UPDATE files SET version=? WHERE id=?", (version, fid))
        c.execute(
            "UPDATE office_checkouts SET base_version=?,digest=? WHERE id=?",
            (version, digest, row["id"]),
        )
        audit(c, f["project_id"], user["id"], "Office-versie opgeslagen", f["name"])
    return {"version": version, "message": f"Versie {version} opgeslagen."}


class CompanyInput(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    primary: str = Field(pattern=r"^#[0-9a-fA-F]{6}$")
    accent: str = Field(pattern=r"^#[0-9a-fA-F]{6}$")
    font: str = Field(min_length=1, max_length=60)
    tagline: str = Field(max_length=200)
    folders: list[str] = Field(min_length=1, max_length=20)


class AISettingsInput(BaseModel):
    provider: str = Field(pattern=r"^(codex|demo)$")
    chat_model: str = Field(pattern=r"^[a-zA-Z0-9][a-zA-Z0-9._-]{0,99}$")
    document_model: str = Field(pattern=r"^[a-zA-Z0-9][a-zA-Z0-9._-]{0,99}$")
    reasoning_effort: str = Field(pattern=r"^(low|medium|high|xhigh)$")


@app.put("/api/settings/ai")
def configure_ai(body: AISettingsInput, user=Depends(current_user)):
    if user["role"] != "admin":
        raise HTTPException(403, "Alleen beheerders mogen AI-instellingen wijzigen")
    with database() as c:
        c.execute(
            "INSERT INTO settings VALUES('ai', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (body.model_dump_json(),),
        )
    return body.model_dump()


@app.post("/api/settings/ai/check")
def check_ai(user=Depends(current_user)):
    if user["role"] != "admin":
        raise HTTPException(403, "Alleen beheerders mogen de verbinding testen")
    if ai.codex.configuration()["provider"] != "codex":
        raise HTTPException(422, "Selecteer en sla eerst Codex op")
    try:
        ai.codex.generate("Antwoord met: Verbinding werkt.", ai.codex.Answer)
    except ValueError as exc:
        raise HTTPException(502, str(exc)) from exc
    return {"message": "Codex-verbinding en chatmodel werken."}


@app.put("/api/settings/company")
def update_company(body: CompanyInput, user=Depends(current_user)):
    admin(user)
    for folder in body.folders:
        clean_name(folder)
    if len({f.lower() for f in body.folders}) != len(body.folders):
        raise HTTPException(422, "Mapnamen moeten uniek zijn")
    with database() as c:
        c.execute("UPDATE settings SET value=? WHERE key='company'", (body.model_dump_json(),))
    return {"ok": True}


@app.patch("/api/users/{target}/role")
def update_role(target: str, body: dict, user=Depends(current_user)):
    admin(user)
    role = body.get("role")
    if role not in {"admin", "consultant"}:
        raise HTTPException(422, "Onbekende rol")
    with database() as c:
        row = c.execute("SELECT * FROM users WHERE id=?", (target,)).fetchone()
        if not row:
            raise HTTPException(404, "Gebruiker niet gevonden")
        if (
            row["role"] == "admin"
            and role != "admin"
            and c.execute("SELECT count(*) FROM users WHERE role='admin'").fetchone()[0] == 1
        ):
            raise HTTPException(422, "Er moet minimaal één beheerder blijven")
        c.execute("UPDATE users SET role=? WHERE id=?", (role, target))
    return {"ok": True}


def seed():
    with database() as c:
        if c.execute("SELECT 1 FROM users").fetchone():
            return
        for user_id, email, name, role in [
            ("demo-admin", "admin@meridian.demo", "Luca van Dijk", "admin"),
            ("demo-owner", "eigenaar@meridian.demo", "Sophie de Vries", "consultant"),
            ("demo-member", "consultant@meridian.demo", "Daan Bakker", "consultant"),
            ("demo-outsider", "nieuw@meridian.demo", "Noor Jansen", "consultant"),
        ]:
            c.execute(
                "INSERT INTO users VALUES(?,?,?,?,?,1)",
                (user_id, email, name, password_hash("MeridianDemo!2026"), role),
            )
        customers = [
            (
                "klant-noord",
                "Noordlicht Energie",
                "Energie",
                "Fictieve regionale energieleverancier met 240 medewerkers. Wil AI inzetten voor klantvragen en kennisdeling.",
                "Eva Smit · Operations manager",
                "Snellere dienstverlening en minder handmatige administratie",
            ),
            (
                "klant-haven",
                "Havenstad Logistiek",
                "Logistiek",
                "Fictieve logistieke dienstverlener. Onderzoekt documentverwerking met AI.",
                "Mark Visser · Directeur",
                "Betrouwbare verwerking van transportdocumenten",
            ),
        ]
        for row in customers:
            c.execute("INSERT INTO customers VALUES(?,?,?,?,?,?,?)", (*row, now()))
            c.execute("INSERT INTO customer_access VALUES(?,'demo-owner')", (row[0],))
        for pid, cid, name, state, desc in [
            (
                "project-ai",
                "klant-noord",
                "AI-strategie & implementatie",
                "In uitvoering",
                "Ontwerp een uitvoerbare AI-roadmap voor de klantenservice. Werk uitsluitend met sampledata.",
            ),
            (
                "project-docs",
                "klant-haven",
                "Slimme documentverwerking",
                "Gepland",
                "Verken automatische classificatie van transportdocumenten.",
            ),
        ]:
            c.execute(
                "INSERT INTO projects VALUES(?,?,?,?,?,'demo-owner',?,'2026-09-14','2026-11-20',0,?)",
                (
                    pid,
                    cid,
                    name,
                    desc,
                    "Een onderbouwd advies, pilotaanpak en meetbare evaluatiecriteria",
                    state,
                    now(),
                ),
            )
            for user_id in ["demo-owner", "demo-member"] if pid == "project-ai" else ["demo-owner"]:
                c.execute("INSERT INTO members VALUES(?,?)", (pid, user_id))
            folder_ids = []
            for folder in company(c)["folders"]:
                fid = uid()
                folder_ids.append(fid)
                c.execute(
                    "INSERT INTO files VALUES(?,?,NULL,?,'folder',1,0,?)",
                    (fid, pid, folder, now()),
                )
            if pid == "project-ai":
                payload = {
                    "type": "docx",
                    "title": "Projectbrief Noordlicht Energie",
                    "sections": [
                        {
                            "heading": "Aanleiding",
                            "body": "Noordlicht Energie onderzoekt AI voor ondersteuning van klantadviseurs. Alle gegevens in deze werkruimte zijn fictief.",
                        },
                        {
                            "heading": "Doel",
                            "body": "Maak een AI-roadmap, selecteer twee pilotprocessen en definieer succescriteria.",
                        },
                        {
                            "heading": "Randvoorwaarden",
                            "body": "Menselijke beoordeling blijft onderdeel van de pilot. Gebruik sampledata en leg aannames vast.",
                        },
                    ],
                }
                save_file(
                    c,
                    pid,
                    folder_ids[0],
                    "Projectbrief.docx",
                    generate(payload, company(c)),
                    "demo-owner",
                )
                payload = {
                    "type": "xlsx",
                    "title": "Sample procesanalyse",
                    "rows": [
                        ["Proces", "Uren per maand", "Pilotgeschikt"],
                        ["Klantvragen", 180, "Ja"],
                        ["Rapportages", 95, "Ja"],
                        ["Kennisbeheer", 120, "Ja"],
                        ["Administratie", 75, "Nee"],
                    ],
                }
                save_file(
                    c,
                    pid,
                    folder_ids[3],
                    "Procesanalyse.xlsx",
                    generate(payload, company(c)),
                    "demo-owner",
                )
                payload = {
                    "type": "pptx",
                    "title": "AI-roadmap",
                    "slides": [
                        {
                            "title": "AI met een helder doel",
                            "body": "Noordlicht Energie\nVerkenning en pilotaanpak",
                            "notes": "Alle gegevens zijn fictief",
                        },
                        {
                            "title": "Begin bij de werkprocessen",
                            "body": "Klantvragen\nKennisbeheer\nRapportages",
                            "notes": "Bespreek de prioriteiten",
                        },
                    ],
                }
                save_file(
                    c,
                    pid,
                    folder_ids[4],
                    "AI-roadmap.pptx",
                    generate(payload, company(c)),
                    "demo-owner",
                )
                for name, instruction, tools in [
                    (
                        "Projectadviseur",
                        "Geef praktisch advies voor AI-implementatie. Benoem aannames en risico's.",
                        ["chat", "documents"],
                    ),
                    (
                        "Data-analist",
                        "Onderbouw analyses met waarden en celverwijzingen. Verzin geen data.",
                        ["chat", "charts", "documents"],
                    ),
                    (
                        "Documentreviewer",
                        "Controleer Nederlandse spelling, grammatica, samenstelling en consistente terminologie.",
                        ["chat", "review"],
                    ),
                ]:
                    c.execute(
                        "INSERT INTO agents VALUES(?,?,?,?,?,'[]')",
                        (uid(), pid, name, instruction, json.dumps(tools)),
                    )
                c.execute(
                    "INSERT INTO messages VALUES(?,?,?,?,?)",
                    (
                        uid(),
                        pid,
                        "demo-owner",
                        "Welkom in het project. De projectbrief en sample procesanalyse staan klaar. Laten we eerst de pilotprocessen afstemmen.",
                        now(),
                    ),
                )
            audit(c, pid, "demo-owner", "Project aangemaakt", "Sampleproject")
